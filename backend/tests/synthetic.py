import io
import zipfile

from docx import Document

SYNTHETIC_SECTIONS = ("SUMMARY", "Experience", "Skills:", "Education")


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]]) -> bytes:
    count = len(pages)
    kids = " ".join(f"{4 + 2 * index} 0 R" for index in range(count))
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {count} >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, lines in enumerate(pages):
        stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(
            f"({_escape(line)}) Tj T*" for line in lines
        )
        stream += " ET"
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {5 + 2 * index} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    body = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, content in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{number} 0 obj\n{content}\nendobj\n".encode("latin-1")
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii")
    for offset in offsets:
        body += f"{offset:010d} 00000 n \n".encode("ascii")
    body += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode("ascii")
    return bytes(body)


def synthetic_resume_lines(tag: str) -> list[str]:
    return [
        f"Sample Person {tag}",
        "SUMMARY",
        "Synthetic engineer used only for automated tests.",
        "Experience",
        "Example Corp Alpha, Senior Builder",
        "Built sample widgets for fictional customers.",
        "Skills:",
        "Widgets, gadgets, testing",
        "Education",
        "Example University, BSc Synthetic Studies",
    ]


def synthetic_pdf(tag: str) -> bytes:
    return make_pdf([synthetic_resume_lines(tag)])


def synthetic_docx(tag: str) -> bytes:
    document = Document()
    document.add_paragraph(f"Sample Person {tag}")
    document.add_heading("Summary", level=1)
    document.add_paragraph("Synthetic engineer used only for automated tests.")
    document.add_heading("Work History", level=1)
    document.add_paragraph("Example Corp Alpha, Senior Builder")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Widgets"
    table.rows[0].cells[1].text = "Gadgets"
    document.add_heading("Education", level=1)
    document.add_paragraph("Example University, BSc Synthetic Studies")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def zip_with(members: dict[str, bytes], *, compress: bool = True) -> bytes:
    buffer = io.BytesIO()
    method = zipfile.ZIP_DEFLATED if compress else zipfile.ZIP_STORED
    with zipfile.ZipFile(buffer, "w", method) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return buffer.getvalue()
